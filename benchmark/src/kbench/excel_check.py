"""Excel fidelity check: the API's xlsx export vs the client's sample BOQ format (Sprint 3).

``sample_spec_from_xls(xls)`` reads the reference layout from the client's legacy ``.xls`` (sheet "כמויות"): title
text, header columns, chapter headers with their numbered notes, and the units per code. ``check_export(xlsx, spec)``
returns a checklist of ``{id, item, status: pass|fail|warn, detail}``. Nothing here prices anything; prices stay empty.

Checks (docs/agents/SPRINT3_BRIEF.md "Excel export"):
X01 main sheet exists · X02 RTL · X03 title rows (title + project name) · X04 header columns in the sample order
(extra columns "רמת ביטחון"/"סטטוס" allowed only at the far left/right) · X05 chapter header rows for every chapter
with lines · X06 chapter notes ("הערות:" + the sample's numbered notes) · X07 codes stored as TEXT · X08 trailing
zeros kept ("4.10", "4.20", "4.00") · X09 code order natural within a chapter · X10 units as in the sample · X11 price
column empty · X12 line total = formula · X13 per-chapter "סה"כ" row with a SUM formula over its lines · X14 grand totals
before/after VAT as formulas · X15 currency number format on totals · X16 extra sheets "פירוט לפי קומה" and "ראיות" ·
X17 no formula-injection text (a text cell starting with = + - @) · X18 quantities equal the BOQ lines file (optional).
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

HEADER = ["סעיף", "תאור העבודה", "יחידה", 'כמות סה"כ', "מחיר יח'", 'סה"כ ₪']
EXTRA_COLS = {"רמת ביטחון", "סטטוס"}
CHAPTER_RE = re.compile(r"^\s*פרק\s*0*(\d+)")
NOTE_RE = re.compile(r"^\s*(\d+)\.\s*")
CODE_RE = re.compile(r"^\d+\.\d+$")


def norm(s: Any) -> str:
    return re.sub(r"\s+", " ", str(s if s is not None else "").replace("\xa0", " ")).strip()


def norm_header(s: Any) -> str:
    return norm(s).replace("״", '"').replace("׳", "'")


def code_text(v: float) -> str:
    return f"{v:.2f}"


# ----------------------------------------------------------------------------- sample


def sample_spec_from_xls(xls: str | Path, sheet: str = "כמויות") -> dict[str, Any]:
    import xlrd

    book = xlrd.open_workbook(str(xls), formatting_info=True)
    s = book.sheet_by_name(sheet)
    title_rows: list[str] = []
    header_row = None
    chapters: dict[str, dict[str, Any]] = {}
    units: dict[str, str] = {}
    cur: str | None = None
    for r in range(s.nrows):
        vals = [s.cell_value(r, c) for c in range(min(6, s.ncols))]
        texts = [norm_header(v) for v in vals]
        if header_row is None:
            if texts[: len(HEADER)] == HEADER:
                header_row = r
            else:
                title_rows += [t for t in texts if t]
            continue
        m = CHAPTER_RE.match(texts[1])
        if m and not texts[0]:
            cur = texts[1]
            chapters[cur] = {"number": m.group(1), "notes": [], "codes": []}
            continue
        if cur and not texts[0] and NOTE_RE.match(texts[1]):
            chapters[cur]["notes"].append(texts[1])
            continue
        if cur and isinstance(vals[0], float):
            code = code_text(vals[0])
            chapters[cur]["codes"].append(code)
            units[code] = texts[2]
    return {
        "sheet": sheet,
        "title_rows": title_rows,
        "header": HEADER,
        "rtl": bool(getattr(s, "columns_from_right_to_left", 0)),
        "chapters": chapters,
        "units": units,
    }


# ----------------------------------------------------------------------------- export


def _code_key(code: str) -> tuple[int, ...]:
    return tuple(int(x) for x in code.split(".") if x.isdigit())


def check_export(xlsx: str | Path, spec: dict[str, Any], lines: list[dict[str, Any]] | None = None) -> list[dict[str, Any]]:
    import openpyxl

    out: list[dict[str, Any]] = []

    def add(cid: str, item: str, ok: bool | None, detail: str = "") -> None:
        out.append({"id": cid, "item": item, "status": "pass" if ok else ("warn" if ok is None else "fail"), "detail": detail})

    wb = openpyxl.load_workbook(str(xlsx), data_only=False)
    ws = wb[spec["sheet"]] if spec["sheet"] in wb.sheetnames else wb.worksheets[0]
    add("X01", f'main sheet "{spec["sheet"]}"', spec["sheet"] in wb.sheetnames, f"sheets: {wb.sheetnames}")
    add("X02", "sheet is right-to-left", bool(ws.sheet_view.rightToLeft), f"rightToLeft={ws.sheet_view.rightToLeft}")

    rows: list[list[Any]] = [list(r) for r in ws.iter_rows()]
    header_idx = None
    col: dict[str, int] = {}
    for i, r in enumerate(rows[:30]):
        texts = [norm_header(c.value) for c in r]
        if all(h in texts for h in HEADER):
            header_idx = i
            col = {h: texts.index(h) for h in HEADER}
            hdr_texts = [t for t in texts if t]
            break
    top = " ".join(norm(c.value) for r in rows[: header_idx or 10] for c in r if c.value is not None)
    want_title = [norm(t).split(" שם הפרויקט")[0] for t in spec["title_rows"] if "כתב כמויות" in str(t)]
    title_ok = bool(want_title) and want_title[0] in top
    add("X03", "title rows (title + project name)", title_ok and "שם הפרויקט" in top, f"top text: {top[:160]!r}")
    if header_idx is None:
        add("X04", "header columns as the sample", False, "no row contains all six sample headers")
        return out
    order = [col[h] for h in HEADER]
    contiguous = order == list(range(order[0], order[0] + len(HEADER)))
    extras = [t for t in hdr_texts if t not in HEADER]
    extras_ok = all(t in EXTRA_COLS for t in extras)
    extras_pos = all(
        (i < order[0] or i > order[-1]) for i, c in enumerate(rows[header_idx]) if norm_header(c.value) in EXTRA_COLS
    )
    add(
        "X04",
        "header columns in sample order; extra columns only at the edges",
        contiguous and extras_ok and extras_pos,
        f"header: {hdr_texts}",
    )

    c_code, c_desc, c_unit, c_qty, c_price, c_tot = order
    chapters_found: dict[str, dict[str, Any]] = {}
    cur: str | None = None
    line_rows: list[tuple[int, list[Any]]] = []
    total_rows: list[tuple[int, list[Any]]] = []
    grand: dict[str, Any] = {}
    for i in range(header_idx + 1, len(rows)):
        r = rows[i]

        def v(ci: int, r: list[Any] = r) -> Any:
            return r[ci].value if ci < len(r) else None

        desc = norm(v(c_desc))
        code_v = v(c_code)
        m = CHAPTER_RE.match(desc)
        if m and code_v in (None, ""):
            cur = desc
            chapters_found[cur] = {"number": m.group(1), "notes": [], "lines": [], "total_row": None, "row": i}
            continue
        if desc.startswith('סה"כ כללי'):
            grand["incl" if "כולל מע" in desc and "לא כולל" not in desc else "excl"] = r[c_tot]
            continue
        if desc == 'סה"כ' and cur:
            chapters_found[cur]["total_row"] = (i, r)
            total_rows.append((i, r))
            continue
        if cur and code_v in (None, "") and (NOTE_RE.match(desc) or desc.startswith("הערות")):
            chapters_found[cur]["notes"].append(desc)
            continue
        if cur and code_v not in (None, ""):
            chapters_found[cur]["lines"].append((i, r))
            line_rows.append((i, r))

    used = [c for c, d in chapters_found.items() if d["lines"]]
    by_num = {d["number"]: c for c, d in spec["chapters"].items()}
    missing_hdr = [c for c in used if chapters_found[c]["number"] not in by_num]
    add(
        "X05",
        "chapter header rows (sample wording) for every chapter with lines",
        bool(used) and not missing_hdr and all(c in spec["chapters"] for c in used),
        f"chapters: {used}; not in sample: {[c for c in used if c not in spec['chapters']]}",
    )

    note_issues = []
    for c in used:
        want = spec["chapters"].get(c) or spec["chapters"].get(by_num.get(chapters_found[c]["number"], ""), {})
        have = {norm(n) for n in chapters_found[c]["notes"]}
        miss = [n for n in want.get("notes", []) if norm(n) not in have]
        if miss:
            note_issues.append(f"{c}: missing {len(miss)}/{len(want.get('notes', []))} notes")
        if want.get("notes") and not any(n.startswith("הערות") for n in have):
            note_issues.append(f"{c}: no 'הערות:' row")
    add("X06", "chapter notes ('הערות:' + numbered notes of the sample)", not note_issues, "; ".join(note_issues))

    codes = [(i, r[c_code]) for i, r in line_rows]
    non_text = [f"r{i + 1}={c.value!r}" for i, c in codes if not isinstance(c.value, str)]
    add("X07", "codes stored as TEXT", not non_text and bool(codes), f"non-text: {non_text[:8]}")
    code_strs = [str(c.value) for _, c in codes]
    tz_expected = [c for c in code_strs if c in {"4.10", "4.20", "4.00"} or (CODE_RE.match(c) and c.endswith("0"))]
    bad_tz = [c for c in code_strs if CODE_RE.match(c) and len(c.split(".")[1]) != 2]
    add(
        "X08",
        "trailing zeros kept (two decimals: 4.10, 4.20, 4.00)",
        not bad_tz,
        f"with trailing zero: {tz_expected}; malformed: {bad_tz}",
    )
    order_issues = []
    for c in used:
        cs = [str(r[c_code].value) for _, r in chapters_found[c]["lines"]]
        if [x for x in cs if CODE_RE.match(x)] != sorted([x for x in cs if CODE_RE.match(x)], key=_code_key):
            order_issues.append(c)
    add("X09", "natural code order inside each chapter (4.09 < 4.10)", not order_issues, f"out of order: {order_issues}")

    unit_issues = [
        f"{r[c_code].value}: {norm(r[c_unit].value)!r} vs {spec['units'][str(r[c_code].value)]!r}"
        for _, r in line_rows
        if str(r[c_code].value) in spec["units"] and norm(r[c_unit].value) != spec["units"][str(r[c_code].value)]
    ]
    add("X10", "units as in the sample", not unit_issues, "; ".join(unit_issues[:8]))
    priced = [str(r[c_code].value) for _, r in line_rows if r[c_price].value not in (None, "")]
    add("X11", "unit-price column empty (MVP)", not priced, f"priced: {priced[:8]}")
    no_formula = [
        str(r[c_code].value) for _, r in line_rows if not (isinstance(r[c_tot].value, str) and r[c_tot].value.startswith("="))
    ]
    add("X12", "line total is a formula (qty × price)", not no_formula, f"no formula: {no_formula[:8]}")
    tot_issues = []
    for c in used:
        tr = chapters_found[c]["total_row"]
        if not tr:
            tot_issues.append(f'{c}: no סה"כ row')
            continue
        f = tr[1][c_tot].value
        if not (isinstance(f, str) and f.upper().startswith("=SUM")):
            tot_issues.append(f"{c}: total {f!r}")
    add("X13", 'per-chapter "סה"כ" row with SUM formula', not tot_issues, "; ".join(tot_issues))
    g_vals = {k: getattr(grand.get(k), "value", None) for k in ("excl", "incl")}
    g_ok = all(isinstance(v, str) and v.startswith("=") for v in g_vals.values())
    add(
        "X14",
        'grand totals "סה"כ כללי לא כולל / כולל מע"מ" as formulas',
        g_ok,
        f"excl={getattr(grand.get('excl'), 'value', None)!r} incl={getattr(grand.get('incl'), 'value', None)!r}",
    )
    fmt_cells = [r[c_tot] for _, r in line_rows + total_rows] + [c for c in grand.values()]
    no_cur = sum(1 for c in fmt_cells if "₪" not in str(c.number_format))
    add("X15", "currency format (₪) on total cells", no_cur == 0 if fmt_cells else None, f"{no_cur}/{len(fmt_cells)} without ₪")
    add(
        "X16",
        'extra sheets "פירוט לפי קומה" and "ראיות"',
        {"פירוט לפי קומה", "ראיות"} <= set(wb.sheetnames),
        f"sheets: {wb.sheetnames}",
    )
    inj = []
    for wsx in wb.worksheets:
        for xrow in wsx.iter_rows():
            for xc in xrow:
                if xc.data_type == "s" and isinstance(xc.value, str) and _injectable(xc.value):
                    inj.append(f"{wsx.title}!{xc.coordinate}")
    add("X17", "no formula-injection text cells", not inj, f"{inj[:8]}")
    if lines is not None:
        q_want: dict[str, float] = {}
        for ln in lines:
            q = ln.get("qtyOverride") if ln.get("qtyOverride") is not None else ln.get("qty", 0)
            q_want[str(ln["code"])] = float(q or 0)
        q_have: dict[str, Any] = {str(r[c_code].value): r[c_qty].value for _, r in line_rows}
        diff = [
            f"{k}: {q_have.get(k)!r} vs {v:g}"
            for k, v in q_want.items()
            if not isinstance(q_have.get(k), (int, float)) or abs(float(q_have[k]) - v) > 1e-9
        ]
        add("X18", "quantities equal the BOQ lines file", not diff, "; ".join(diff[:10]))
    return out


def _injectable(v: str) -> bool:
    """A text cell Excel would evaluate: leading = + @, or - not followed by a space (SEC formula injection)."""
    return v[:1] in ("=", "+", "@") or (v[:1] == "-" and v[1:2] not in (" ", ""))


def render_checklist(items: list[dict[str, Any]], title: str = "Excel fidelity check") -> str:
    mark = {"pass": "[x]", "fail": "[ ]", "warn": "[~]"}
    lines = [f"# {title}", ""]
    for it in items:
        d = f" — {it['detail']}" if it["detail"] and it["status"] != "pass" else ""
        lines.append(f"- {mark[it['status']]} {it['id']} {it['item']}{d}")
    n = {s: sum(1 for it in items if it["status"] == s) for s in ("pass", "fail", "warn")}
    lines += ["", f"{n['pass']} pass, {n['fail']} fail, {n['warn']} warn"]
    return "\n".join(lines) + "\n"
