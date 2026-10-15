"""Excel fidelity checklist on synthetic workbooks (no customer content)."""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from kbench.excel_check import _injectable, check_export, render_checklist, sample_spec_from_xls

CH = "פרק 04 – עבודות חשמל"
SPEC: dict[str, Any] = {
    "sheet": "כמויות",
    "title_rows": ["כתב כמויות ומחירים", "כתב כמויות ומחירים שם הפרויקט: בית לדוגמה"],
    "rtl": True,
    "chapters": {CH: {"number": "4", "notes": ["1. הערה ראשונה", "2. הערה שנייה"], "codes": ["4.09", "4.10"]}},
    "units": {"4.09": "יחידה", "4.10": "יחידה"},
}
HDR = ["סעיף", "תאור העבודה", "יחידה", 'כמות סה"כ', "מחיר יח'", 'סה"כ ₪']


def _book(
    path: Path,
    *,
    codes: tuple[Any, Any] = ("4.09", "4.10"),
    rtl: bool = True,
    notes: bool = True,
    extra_left: bool = True,
    inject: bool = False,
    sheets: bool = True,
) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "כמויות"
    ws.sheet_view.rightToLeft = rtl
    pre = ["רמת ביטחון"] if extra_left else []
    off = len(pre)
    ws.append(["כתב כמויות ומחירים"])
    ws.append(["כתב כמויות ומחירים\nשם הפרויקט: בית לדוגמה"])
    ws.append(pre + HDR + ["סטטוס"])
    ws.append([None] * off + [None, CH])
    if notes:
        ws.append([None] * off + [None, "הערות:"])
        ws.append([None] * off + [None, "1.\xa0\xa0 הערה ראשונה"])
        ws.append([None] * off + [None, "2. הערה שנייה"])
    first = ws.max_row + 1
    for code, q in zip(codes, (16, 2), strict=True):
        r = ws.max_row + 1
        d, e = openpyxl.utils.get_column_letter(off + 4), openpyxl.utils.get_column_letter(off + 5)
        ws.append(
            (["מאומת"] if extra_left else [])
            + [code, "=cmd" if inject else "תיאור", "יחידה", q, None, f"={d}{r}*{e}{r}", "טיוטה"]
        )
        ws.cell(r, off + 6).number_format = '"₪"\\ #,##0'
    last = ws.max_row
    f = openpyxl.utils.get_column_letter(off + 6)
    ws.append([None] * off + [None, 'סה"כ', None, None, None, f"=SUM({f}{first}:{f}{last})"])
    ws.cell(ws.max_row, off + 6).number_format = '"₪"\\ #,##0'
    tr = ws.max_row
    ws.append([None] * off + [None, 'סה"כ כללי לא כולל מע"מ', None, None, None, f"={f}{tr}"])
    ws.cell(ws.max_row, off + 6).number_format = '"₪"\\ #,##0'
    ws.append([None] * off + [None, 'סה"כ כללי כולל מע"מ', None, None, None, f"={f}{ws.max_row}*1.18"])
    ws.cell(ws.max_row, off + 6).number_format = '"₪"\\ #,##0'
    if inject:
        ws.cell(first, off + 2).data_type = "s"
    if sheets:
        wb.create_sheet("פירוט לפי קומה")
        wb.create_sheet("ראיות")
    wb.save(path)
    return path


def _status(items: list[dict[str, Any]]) -> dict[str, str]:
    return {it["id"]: it["status"] for it in items}


def test_good_export_passes(tmp_path: Path) -> None:
    items = check_export(
        _book(tmp_path / "a.xlsx"), SPEC, [{"code": "4.09", "qty": 16}, {"code": "4.10", "qty": 1, "qtyOverride": 2}]
    )
    st = _status(items)
    assert all(v == "pass" for v in st.values()), [it for it in items if it["status"] != "pass"]
    assert "X18" in st
    assert "0 fail" in render_checklist(items)


def test_numeric_codes_rtl_notes_sheets(tmp_path: Path) -> None:
    st = _status(check_export(_book(tmp_path / "b.xlsx", codes=(4.09, 4.1), rtl=False, notes=False, sheets=False), SPEC))
    assert st["X02"] == "fail" and st["X06"] == "fail" and st["X07"] == "fail" and st["X16"] == "fail"


def test_order_and_trailing_zero(tmp_path: Path) -> None:
    st = _status(check_export(_book(tmp_path / "c.xlsx", codes=("4.10", "4.1")), SPEC))
    assert st["X08"] == "fail" and st["X09"] == "fail"


def test_quantity_mismatch(tmp_path: Path) -> None:
    items = check_export(_book(tmp_path / "d.xlsx"), SPEC, [{"code": "4.09", "qty": 15}])
    assert _status(items)["X18"] == "fail"


def test_injection_helper() -> None:
    assert _injectable("=1+1") and _injectable("@x") and _injectable("+972") and _injectable("-2")
    assert not _injectable("- note") and not _injectable("text") and not _injectable("")


@pytest.mark.skipif(
    not (Path(os.environ.get("KAMUTAN_DATA_DIR", "/home/claude/kamutan-data")) / "sample-noah" / "boq_sample.xls").exists(),
    reason="customer sample not available",
)
def test_sample_spec_from_real_xls() -> None:
    spec = sample_spec_from_xls(
        Path(os.environ.get("KAMUTAN_DATA_DIR", "/home/claude/kamutan-data")) / "sample-noah" / "boq_sample.xls"
    )
    assert spec["rtl"] is True
    ch4 = next(v for k, v in spec["chapters"].items() if v["number"] == "4")
    assert "4.10" in ch4["codes"] and "4.00" in ch4["codes"] and "4.20" in ch4["codes"]


def test_injection_cell_flagged(tmp_path: Path) -> None:
    st = _status(check_export(_book(tmp_path / "e.xlsx", inject=True), SPEC))
    assert st["X17"] == "fail"
