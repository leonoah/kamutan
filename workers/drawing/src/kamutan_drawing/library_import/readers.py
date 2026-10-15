"""Spreadsheet readers for past BOQ files (``.xls`` via xlrd, ``.xlsx`` via openpyxl).

Runs **inside the parser sandbox** only (task ``boq_parse``): both libraries parse hostile
binary/zip input. Output is a plain, bounded grid per sheet::

    Sheet(name, rows=[[Cell(value, kind, decimals), ...], ...])

``kind`` is ``"text"`` | ``"number"`` | ``"empty"``; ``decimals`` is the number of decimals the
cell's number format shows (``"0.00"`` → 2, ``General`` → None) — the code column of a BOQ is
usually a number formatted ``0.00`` (4.1 displayed as "4.10"), which is how TEXT codes are restored.
Formulas are read as their cached values (xlsx ``data_only``), never evaluated.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Any

MAX_SHEETS = 32
MAX_ROWS = 5000
MAX_COLS = 40
MAX_TEXT = 4000

XLS_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ZIP_MAGIC = b"PK\x03\x04"


@dataclass(frozen=True)
class Cell:
    value: str | float | None
    kind: str  # text | number | empty
    decimals: int | None = None


EMPTY = Cell(None, "empty")


@dataclass
class Sheet:
    name: str
    rows: list[list[Cell]] = field(default_factory=list)

    def to_obj(self) -> dict[str, Any]:
        return {"name": self.name, "rows": [[[c.value, c.kind, c.decimals] for c in r] for r in self.rows]}

    @staticmethod
    def from_obj(o: dict[str, Any]) -> Sheet:
        return Sheet(str(o["name"]), [[Cell(v, str(k), d) for v, k, d in r] for r in o["rows"]])


def sniff(head: bytes) -> str | None:
    """``xls`` (OLE2 compound file) | ``xlsx`` (zip) | None."""
    if head.startswith(XLS_MAGIC):
        return "xls"
    if head.startswith(ZIP_MAGIC):
        return "xlsx"
    return None


_DEC_RE = re.compile(r"0\.(0+)")


def decimals_of(fmt: str | None) -> int | None:
    """Decimals shown by an Excel number format (first section). ``General``/text → None."""
    if not fmt:
        return None
    first = fmt.split(";", 1)[0]
    first = re.sub(r'"[^"]*"|\[[^\]]*\]', "", first)  # literals, colours/conditions
    m = _DEC_RE.search(first)
    if m:
        return len(m.group(1))
    if re.search(r"(^|[^.#0])0($|[^.0])", first) or first.strip() in ("0", "#,##0", "#"):
        return 0
    return None


def _text(v: Any) -> Cell:
    s = str(v)[:MAX_TEXT]
    return Cell(s, "text") if s.strip() else EMPTY


def read_xls(path: str) -> list[Sheet]:
    import xlrd

    book = xlrd.open_workbook(path, formatting_info=True, on_demand=True)
    out: list[Sheet] = []
    try:
        for si in range(min(book.nsheets, MAX_SHEETS)):
            sh = book.sheet_by_index(si)
            rows: list[list[Cell]] = []
            for r in range(min(sh.nrows, MAX_ROWS)):
                row: list[Cell] = []
                for c in range(min(sh.ncols, MAX_COLS)):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_TEXT:
                        row.append(_text(cell.value))
                    elif cell.ctype == xlrd.XL_CELL_NUMBER:
                        dec: int | None = None
                        try:
                            xf = book.xf_list[cell.xf_index]
                            dec = decimals_of(book.format_map[xf.format_key].format_str)
                        except (IndexError, KeyError, AttributeError, TypeError):
                            dec = None
                        row.append(Cell(float(cell.value), "number", dec))
                    else:
                        row.append(EMPTY)
                rows.append(row)
            out.append(Sheet(str(sh.name)[:100], rows))
            book.unload_sheet(si)
    finally:
        book.release_resources()
    return out


def read_xlsx(path: str) -> list[Sheet]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    out: list[Sheet] = []
    try:
        for ws in wb.worksheets[:MAX_SHEETS]:
            rows: list[list[Cell]] = []
            for r_i, r in enumerate(ws.iter_rows(max_col=MAX_COLS)):
                if r_i >= MAX_ROWS:
                    break
                row: list[Cell] = []
                for cell in r:
                    v = getattr(cell, "value", None)
                    if v is None:
                        row.append(EMPTY)
                    elif isinstance(v, bool):
                        row.append(_text(str(v)))
                    elif isinstance(v, int | float):
                        row.append(Cell(float(v), "number", decimals_of(getattr(cell, "number_format", None))))
                    else:
                        row.append(_text(v))
                rows.append(row)
            out.append(Sheet(str(ws.title)[:100], rows))
    finally:
        wb.close()
    return out


def read_workbook(path: str, kind: str) -> list[Sheet]:
    if kind == "xls":
        return read_xls(path)
    if kind == "xlsx":
        return read_xlsx(path)
    raise ValueError(f"unsupported workbook kind {kind}")
