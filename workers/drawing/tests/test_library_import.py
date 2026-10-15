"""library.import: workbook readers, BOQ structure parsing (TEXT code restoration), Firestore upsert, job."""

from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path

import pytest

from kamutan_drawing.boq.contracts import LibraryImportJob
from kamutan_drawing.errors import IngestError
from kamutan_drawing.library_import.importer import (
    DEFAULT_TEMPLATE_ID,
    DEFAULT_TEMPLATE_NAME,
    item_id,
    run_library_import,
    write_library,
)
from kamutan_drawing.library_import.parse import parse_workbook
from kamutan_drawing.library_import.readers import Cell, Sheet, decimals_of, read_workbook, sniff
from kamutan_drawing.sandbox.runner import Limits
from kamutan_drawing.stores.memory import InMemoryBlobStore, InMemoryDocStore

from .conftest import DATA_DIR

LIMITS = Limits(net="off", landlock="off", drop_root="off")


def synth_xlsx(*, general_codes: bool = False) -> bytes:
    """A BOQ like the client's sample: title rows, header, chapters with notes, numeric codes, totals, extra sheet."""
    import openpyxl

    wb = openpyxl.Workbook()
    cover = wb.active
    assert cover is not None
    cover.title = "כותרת"
    cover["A1"] = "דף שער"
    ws = wb.create_sheet("כמויות")
    ws.append([None, None, None, None, "ישראל ישראלי"])
    ws.append(["כתב כמויות ומחירים\nשם הפרויקט: בית לדוגמה, עיר"])
    ws.append(["סעיף", "תאור העבודה", "יחידה", 'כמות סה"כ', "מחיר יח'", 'סה"כ ₪'])
    ws.append([None, "פרק 04 – עבודות חשמל"])
    ws.append([None, "הערות: "])
    ws.append([None, "1.\xa0\xa0\xa0\xa0 כל העבודות לפי המפרט"])
    ws.append([None, "3.     לכמויות ראה תכניות"])
    rows = [
        (4.0, "שקע בודד", 23),
        (4.01, "שקע כפול", 15),
        (4.09, "נק' מאור כוללת מפסק", 16),
        (4.1, "נק' מאור בעלת שני מפסקים מחליפים", 2),
        (4.11, "שתי נק' מאור משותפות בעלות מפסק בודד", 5),
        (4.12, "הוספת מפסק מחליף לנקודת מאור קיימת", 3),
        (4.14, "הוספת מפסק למאוורר", 1),
        (4.2, "צינור קוברה", 5),
        (4.21, "שקע תלת פאזי", 3),
    ]
    for code, desc, qty in rows:
        ws.append([code, desc, "יחידה", qty, None, 0])
        if not general_codes:
            ws.cell(ws.max_row, 1).number_format = "0.00"
    ws.append([None, 'סה"כ', None, None, None, 0])
    ws.append([None, "פרק 08 - שונות"])
    ws.append(["9.01", "התקנת גופי התאורה בבית", "יחידה", 43, 0, 0])  # TEXT code
    ws.append([None, 'סה"כ', None, None, None, 0])
    ws.append([None, 'סה"כ כללי לא כולל מע"מ', None, None, None, 0])
    ws.append([None, 'סה"כ כללי כולל מע"מ', None, None, None, 0])
    other = wb.create_sheet("שינויים")
    other.append(["סעיף", "תאור", "יח'", "כמות"])
    other.append([1.01, "שינוי", 'מ"ר', 1])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse_bytes(data: bytes, tmp_path: Path, kind: str = "xlsx") -> dict:
    p = tmp_path / f"b.{kind}"
    p.write_bytes(data)
    return parse_workbook(read_workbook(str(p), kind))


def test_item_id_matches_contracts_item_id_for():
    # vectors produced by packages/contracts itemIdFor (TS)
    assert item_id("4.10", "נק' מאור בעלת שני מפסקים מחליפים") == "it_4-10_d7c6d7627f422152"
    assert item_id("4.10", "  נק' מאור\u00a0 בעלת שני מפסקים מחליפים ") == "it_4-10_d7c6d7627f422152"
    assert item_id("A/1", " x\u00a0 y ") == "it_A-1_a71f7e125799f21c"
    assert item_id("4.1", "נק' מאור בעלת שני מפסקים מחליפים") != item_id("4.10", "נק' מאור בעלת שני מפסקים מחליפים")


def test_decimals_of():
    assert decimals_of("0.00") == 2
    assert decimals_of("#,##0.0") == 1
    assert decimals_of("General") is None
    assert decimals_of('0.00" ₪";[Red]-0.00') == 2
    assert decimals_of(None) is None


def test_sniff():
    assert sniff(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1") == "xls"
    assert sniff(b"PK\x03\x04xxxx") == "xlsx"
    assert sniff(b"%PDF-1.7") is None


@pytest.mark.parametrize("general", [False, True])
def test_parse_synthetic_xlsx(tmp_path, general):
    p = _parse_bytes(synth_xlsx(general_codes=general), tmp_path)
    assert p["schema"] == "kamutan.boq-import/1"
    assert p["sheet"] == "כמויות" and "שינויים" in p["ignoredSheets"]
    # project name / person never reach the template
    assert p["titleRows"] == ["כתב כמויות ומחירים", "שם הפרויקט: {projectName}"]
    assert [c["key"] for c in p["columns"]] == ["code", "description", "unit", "qty", "unitPrice", "total"]
    chapters = {c["code"]: c for c in p["chapters"]}
    assert set(chapters) == {"04", "08"}
    assert chapters["04"]["title"] == "עבודות חשמל"
    assert chapters["04"]["notes"] == ["1. כל העבודות לפי המפרט", "3. לכמויות ראה תכניות"]
    items = {i["code"]: i for i in p["items"]}
    # TEXT codes restored: from the number format, or (General) from the sequence 4.09, 4.1, 4.11
    assert {"4.00", "4.01", "4.09", "4.10", "4.11", "4.12", "4.14", "4.20", "4.21", "9.01"} == set(items)
    assert "4.1" not in items and "4.2" not in items
    assert "trailing_zero_restored" in items["4.10"]["flags"]
    assert "ambiguous_code" in items["4.00"]["flags"]  # "4.00" vs chapter number "4"
    assert "gap_before:4.13" in items["4.14"]["flags"]
    assert "chapter_code_mismatch:8" in items["9.01"]["flags"]
    assert items["9.01"]["chapter"] == "08" and items["9.01"]["qty"] == 43
    assert items["4.10"]["unit"] == "יחידה" and items["4.10"]["chapterTitle"] == "עבודות חשמל"
    assert p["totals"]["grand"] == ['סה"כ כללי לא כולל מע"מ', 'סה"כ כללי כולל מע"מ']


def test_sequence_restoration_without_format_is_flagged_when_ambiguous():
    def num(v: float) -> Cell:
        return Cell(v, "number", None)

    def txt(s: str) -> Cell:
        return Cell(s, "text")

    def sheet(*codes: float) -> Sheet:
        rows = [[txt("סעיף"), txt("תאור"), txt("יחידה"), txt("כמות")], [Cell(None, "empty"), txt("פרק 02 - בנייה")]]
        rows += [[num(c), txt(f"פריט {i}"), txt("יח'"), num(1)] for i, c in enumerate(codes)]
        return Sheet("x", rows)

    # 2.09, 2.1, 2.11: the sequence supports "2.10" → restored, not ambiguous
    codes = {i["code"]: i["flags"] for i in parse_workbook([sheet(2.09, 2.1, 2.11)])["items"]}
    assert "trailing_zero_restored" in codes["2.10"] and "ambiguous_code" not in codes["2.10"]
    # 2.01, 2.1, 2.3: neither "2.1" nor "2.10" fits → padded to the file's 2 decimals, flagged
    codes = {i["code"]: i["flags"] for i in parse_workbook([sheet(2.01, 2.1, 2.3)])["items"]}
    assert "ambiguous_code" in codes["2.10"]
    # TEXT codes are kept as written
    p = parse_workbook([Sheet("x", [[txt("סעיף"), txt("תאור"), txt("יחידה")], [txt("4.1"), txt("א"), txt("יח'")]])])
    assert p["items"][0]["code"] == "4.1"


def test_no_boq_sheet():
    with pytest.raises(ValueError):
        parse_workbook([Sheet("x", [[Cell("hello", "text")]])])


def test_write_library_idempotent_and_preserves_manual(tmp_path):
    parsed = _parse_bytes(synth_xlsx(), tmp_path)
    docs = InMemoryDocStore()
    out = write_library("o1", IMP, parsed, docs)
    assert out.created == 10 and out.template_id == DEFAULT_TEMPLATE_ID
    lib = dict(docs.list_docs("orgs/o1/itemLibrary"))
    assert len(lib) == 10
    iid = item_id("4.10", "נק' מאור בעלת שני מפסקים מחליפים")
    assert lib[iid]["code"] == "4.10" and lib[iid]["source"] == "import" and lib[iid]["active"] is True
    assert lib[iid]["sourceRef"] == IMP and lib[iid]["chapter"] == "04" and lib[iid]["tags"] == []
    tpl = docs.get(f"orgs/o1/boqTemplates/{DEFAULT_TEMPLATE_ID}")
    assert tpl["default"] is True and [c["code"] for c in tpl["chapters"]] == ["04", "08"]
    assert tpl["chapters"][0]["title"] == "פרק 04 – עבודות חשמל" and tpl["chapters"][1]["codePrefixes"] == ["9"]
    assert (
        tpl["source"] == "import" and tpl["vatRate"] == 0.18 and [c["key"] for c in tpl["columns"]][:2] == ["code", "description"]
    )
    assert tpl["titleRows"][1] == "שם הפרויקט: {projectName}"
    # user deactivates one item and edits another to manual; re-import changes neither and adds no docs
    docs.set(f"orgs/o1/itemLibrary/{iid}", {"active": False})
    other = item_id("4.01", "שקע כפול")
    docs.set(f"orgs/o1/itemLibrary/{other}", {"source": "manual", "description": "שקע כפול (ערוך)"})
    docs.set(f"orgs/o1/boqTemplates/{DEFAULT_TEMPLATE_ID}", {"name": "שלי"})
    again = write_library("o1", IMP, parsed, docs)
    assert again.created == 0 and again.template_id is None
    lib2 = dict(docs.list_docs("orgs/o1/itemLibrary"))
    assert len(lib2) == 10
    assert lib2[iid]["active"] is False
    assert lib2[other]["description"] == "שקע כפול (ערוך)"
    assert docs.get(f"orgs/o1/boqTemplates/{DEFAULT_TEMPLATE_ID}")["name"] == "שלי"


IMP = "0b3c1f6e-1234-4abc-8def-0123456789ab"


def _job(data: bytes, kind: str = "xlsx", **kw) -> tuple[LibraryImportJob, InMemoryBlobStore, InMemoryDocStore]:
    path = f"orgs/o1/imports/{IMP}.{kind}"
    blobs, docs = InMemoryBlobStore(), InMemoryDocStore()
    blobs.write(path, data, "application/octet-stream")
    docs.set("jobs/j1", {"orgId": "o1", "type": "library.import", "status": "queued", "attempts": 0})
    docs.set(f"orgs/o1/libraryImports/{IMP}", {"status": "queued", "orgId": "o1"})
    job = LibraryImportJob.model_validate(
        {"schemaVersion": 1, "type": "library.import", "jobId": "j1", "orgId": "o1", "projectId": None, "importId": IMP,
         "storagePath": path, "sourceKind": kind, "traceparent": None, **kw}
    )  # fmt: skip
    return job, blobs, docs


def test_import_job_end_to_end_sandboxed():
    job, blobs, docs = _job(synth_xlsx())
    out = run_library_import(job, blobs, docs, limits=LIMITS)
    assert out.status == "succeeded" and out.items == 10
    jd = docs.get("jobs/j1")
    assert jd["status"] == "succeeded" and jd["result"]["importId"] == IMP and jd["result"]["templateId"] == DEFAULT_TEMPLATE_ID
    assert jd["result"]["sha256"] == hashlib.sha256(blobs.read(job.storagePath)).hexdigest()
    imp = docs.get(f"orgs/o1/libraryImports/{IMP}")
    assert imp["status"] == "succeeded" and imp["stats"]["items"] == 10 and imp["stats"]["templates"] == 1
    assert imp["stats"]["ambiguousCodes"] == ["4.00"] and imp["stats"]["templateName"] == DEFAULT_TEMPLATE_NAME
    assert jd["result"]["ambiguousCodes"] == ["4.00"]
    # redelivery: no-op
    assert run_library_import(job, blobs, docs, limits=LIMITS).status == "skipped"
    assert len(docs.list_docs("orgs/o1/itemLibrary")) == 10


def test_import_job_rejects_bad_input():
    job, blobs, docs = _job(b"%PDF-1.7 not a workbook")
    with pytest.raises(IngestError) as ei:
        run_library_import(job, blobs, docs, limits=LIMITS)
    assert ei.value.code == "unsupported_format" and not ei.value.retryable
    assert docs.get("jobs/j1")["status"] == "failed"
    assert docs.get(f"orgs/o1/libraryImports/{IMP}")["error"]["code"] == "unsupported_format"
    # xls declared, xlsx content
    job, blobs, docs = _job(synth_xlsx(), kind="xls")
    with pytest.raises(IngestError) as ei:
        run_library_import(job, blobs, docs, limits=LIMITS)
    assert ei.value.code == "file_type_mismatch"
    # corrupt zip → parsed in the sandbox → typed error, never a crash
    job, blobs, docs = _job(b"PK\x03\x04" + b"\x00" * 200)
    with pytest.raises(IngestError) as ei:
        run_library_import(job, blobs, docs, limits=LIMITS)
    assert ei.value.code == "file_unreadable"
    # foreign org job doc
    job, blobs, docs = _job(synth_xlsx())
    docs.set("jobs/j1", {"orgId": "other"})
    with pytest.raises(IngestError) as ei:
        run_library_import(job, blobs, docs, limits=LIMITS)
    assert ei.value.code == "invalid_payload"


def test_import_job_payload_validation():
    data = synth_xlsx()
    with pytest.raises(ValueError):
        _job(data, storagePath=f"orgs/other/imports/{IMP}.xlsx")
    with pytest.raises(ValueError):
        _job(data, storagePath=f"orgs/o1/imports/{IMP}.xls")
    with pytest.raises(ValueError):
        _job(data, importId="not-a-uuid")


SAMPLE_XLS = DATA_DIR / "boq_sample.xls"
TRUTH = DATA_DIR.parent / "gt" / "noah" / "boq_truth.json"


@pytest.mark.skipif(not SAMPLE_XLS.exists(), reason="sample BOQ not available (set KAMUTAN_DATA_DIR)")
def test_sample_xls_matches_truth(tmp_path):
    p = parse_workbook(read_workbook(str(SAMPLE_XLS), "xls"))
    assert p["sheet"] == "כמויות"
    items = {i["code"]: i for i in p["items"]}
    assert len(items) == 77
    for code in ("4.10", "4.20", "1.10", "4.00", "9.01"):
        assert code in items
    assert items["4.00"]["description"] == "שקע בודד" and "ambiguous_code" in items["4.00"]["flags"]
    assert [c["code"] for c in p["chapters"]] == ["01", "02", "03", "04", "05", "06", "07", "08"]
    assert len({c["code"]: c for c in p["chapters"]}["04"]["notes"]) == 5
    if TRUTH.exists():
        truth = {ln["code"]: ln for ln in json.loads(TRUTH.read_text())["lines"]}
        assert set(truth) == set(items)
        for c, t in truth.items():
            assert items[c]["qty"] == t["quantity"], c
            assert items[c]["unit"] == t["unit"], c
            assert ("ambiguous_code" in t["flags"]) == ("ambiguous_code" in items[c]["flags"]), c
